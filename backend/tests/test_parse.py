import io
from datetime import date, datetime

import openpyxl
import pytest

from app.ingest.parse import parse_csv, parse_tabular_file, parse_xlsx


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    first = wb.active
    first.title = "Data"
    first.append(["id", "day", "ts", "flag", "note"])
    first.append([1, date(2024, 1, 15), datetime(2024, 1, 15, 9, 30), True, None])
    first.append([2, date(2024, 1, 16), datetime(2024, 1, 16, 0, 0), False, "x"])
    second = wb.create_sheet("Other")
    second.append(["ignored"])
    second.append(["also ignored"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestCsv:
    def test_utf8(self):
        parsed = parse_csv(b"id,name\n1,Ana\n")
        assert parsed.headers == ["id", "name"]
        assert parsed.rows == [["1", "Ana"]]

    def test_quoted_commas_and_newlines(self):
        parsed = parse_csv(b'id,note\n1,"a, b"\n2,"line1\nline2"\n')
        assert parsed.rows == [["1", "a, b"], ["2", "line1\nline2"]]

    def test_blank_rows_dropped(self):
        parsed = parse_csv(b"id,v\n\n1,2\n,\n   \n3,4\n")
        assert parsed.rows == [["1", "2"], ["3", "4"]]

    def test_windows_1252_fallback(self):
        parsed = parse_csv("id,name\n1,café\n".encode("cp1252"))
        assert parsed.rows == [["1", "café"]]

    def test_utf8_bom_header_is_not_prefixed(self):
        parsed = parse_csv(b"\xef\xbb\xbfid,name\n1,x\n")
        assert parsed.headers == ["id", "name"]
        assert parsed.headers[0] == "id"


class TestXlsx:
    def test_first_sheet_only(self):
        parsed = parse_xlsx(make_xlsx())
        assert parsed.headers == ["id", "day", "ts", "flag", "note"]
        assert all(row[0] != "ignored" for row in parsed.rows)
        assert len(parsed.rows) == 2

    def test_cell_serialization(self):
        parsed = parse_xlsx(make_xlsx())
        assert parsed.rows[0] == ["1", "2024-01-15", "2024-01-15 09:30:00", "true", ""]
        assert parsed.rows[1] == ["2", "2024-01-16", "2024-01-16", "false", "x"]


class TestXlsxCloseOnError:
    def test_workbook_closed_when_iteration_fails(self, monkeypatch):
        closed: list[bool] = []

        class FailingSheet:
            def iter_rows(self, values_only=False):
                raise RuntimeError("boom")

        class FakeWorkbook:
            worksheets = [FailingSheet()]

            def close(self):
                closed.append(True)

        monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: FakeWorkbook())
        with pytest.raises(RuntimeError, match="boom"):
            parse_xlsx(b"ignored")
        assert closed == [True]


class TestDispatch:
    @pytest.mark.parametrize("name", ["data.csv", "DATA.CSV", "Data.Csv"])
    def test_csv_extension_case_insensitive(self, name):
        parsed = parse_tabular_file(name, b"id\n1\n")
        assert parsed.rows == [["1"]]

    @pytest.mark.parametrize("name", ["data.xlsx", "DATA.XLSX"])
    def test_xlsx_extension_case_insensitive(self, name):
        parsed = parse_tabular_file(name, make_xlsx())
        assert parsed.headers[0] == "id"

    def test_unsupported_extension(self):
        with pytest.raises(ValueError, match="unsupported file type"):
            parse_tabular_file("data.json", b"{}")
